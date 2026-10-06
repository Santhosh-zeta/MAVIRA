"""API + end-to-end: upload -> parse -> criticality -> assessment -> evidence ->
risk -> review -> export, against the real Meridian workbook and the cached
OSINT snapshot. Runs synchronously by patching the dispatcher."""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.evidence.validator import validate_excerpt
from app.services.pipeline import orchestrator


@pytest.fixture(scope="module")
def client():
    import app.api.routes.core as core
    core.dispatch = lambda run_id: (orchestrator.run_assessment(run_id), "sync")[1]
    from app.main import app
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def completed(client, sample_bytes):
    r = client.post("/api/imports", files={"file": ("Meridian_Vendor_Input.xlsx", sample_bytes,
                                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["accepted"] and body["validation"]["vendor_count"] == 6
    run = client.post("/api/runs", json={"import_id": body["import_id"]}).json()
    status = client.get(f"/api/runs/{run['run_id']}").json()
    return {"import_id": body["import_id"], "run": status}


def test_health_reports_no_egress_in_demo(client):
    h = client.get("/api/health").json()
    assert h["osint_mode"] == "DEMO" and not h["llm_enabled"]
    assert any("No network egress" in x for x in h["data_egress"])


def test_upload_rejects_non_workbook(client):
    r = client.post("/api/imports", files={"file": ("evil.xlsx", b"<script>", "application/octet-stream")})
    assert r.status_code == 400 and "error" in r.json()


def test_run_completes_all_vendors(completed):
    run = completed["run"]
    assert run["status"] == "COMPLETED" and run["completed_vendors"] == 6 and run["failed_vendors"] == 0
    assert run["progress"] == 1.0
    assert run["analyzer"].startswith("rule-based") and run["methodology_version"]
    for a in run["assessments"]:
        assert all(s["done"] for s in a["stages"])


def test_expected_outcomes(client, completed):
    by = {a["vendor_ref"]: a for a in completed["run"]["assessments"]}
    assert by["V-002"]["criticality_tier"] == "CRITICAL" and by["V-002"]["ai_usage"] == "YES"
    assert by["V-005"]["ai_usage"] == "YES"
    assert by["V-006"]["ai_usage"] == "INCONCLUSIVE" and by["V-006"]["risk_class"] == "NONE_IDENTIFIED"
    for ref in ("V-001", "V-003", "V-004"):
        assert by[ref]["ai_usage"] == "INCONCLUSIVE"       # signals exist but nothing confirms service AI
    assert by["V-001"]["risk_class"] in ("MEDIUM", "HIGH")   # privileged access + AI-capable platform


def test_every_valid_excerpt_exists_in_its_source(client, completed):
    for a in completed["run"]["assessments"]:
        d = client.get(f"/api/assessments/{a['id']}").json()
        for e in d["evidence"]:
            if e["validation_status"] != "VALID":
                continue
            src = client.get(f"/api/sources/{e['source']['id']}").json()
            assert validate_excerpt(e["exact_excerpt"], src["content"]).valid
            assert src["content"][e["excerpt_start"]:e["excerpt_end"]]
            assert src["collection_method"] == "CACHED" and src["retrieved_at"]


def test_detail_contains_traceability(client, completed):
    a = next(x for x in completed["run"]["assessments"] if x["vendor_ref"] == "V-002")
    d = client.get(f"/api/assessments/{a['id']}").json()
    assert d["criticality_factors"]["factors"]["data_sensitivity"]["score"] == 1.0
    assert d["depth_plan"]["tier"] == "CRITICAL" and d["searches"]
    assert {p["provider_name"] for p in d["providers"]} == {"Microsoft (Azure OpenAI / Foundry / Copilot)"}
    assert all(p["service_linkage"] == "ENTERPRISE_LEVEL_ONLY" for p in d["providers"])
    assert len(d["risk_factors"]) == 9 and d["recommendations"]
    assert d["negative_evidence"]["sources_checked"] >= 1


def test_negative_evidence_recorded_for_tch(client, completed):
    a = next(x for x in completed["run"]["assessments"] if x["vendor_ref"] == "V-006")
    d = client.get(f"/api/assessments/{a['id']}").json()
    ne = d["negative_evidence"]
    assert ne["statement"] == "NO MEANINGFUL PUBLIC EVIDENCE IDENTIFIED WITHIN ASSESSED OSINT SCOPE"
    assert ne["searches_performed"] > 0 and ne["assessment_scope"]


def test_review_flow_preserves_original(client, completed):
    a = next(x for x in completed["run"]["assessments"] if x["vendor_ref"] == "V-004")
    url = f"/api/assessments/{a['id']}/reviews"
    bad = client.post(url, json={"decision": "OVERRIDE", "analyst": "J. Analyst", "overrides": {"risk_class": "MEDIUM"}})
    assert bad.status_code == 400                                      # reason required
    ok = client.post(url, json={"decision": "OVERRIDE", "analyst": "J. Analyst", "reason": "Vendor confirmed AI notetaking on client calls in QBR",
                                "overrides": {"risk_class": "MEDIUM", "ai_usage": "YES"}})
    assert ok.status_code == 200 and ok.json()["review_status"] == "OVERRIDDEN"
    d = client.get(f"/api/assessments/{a['id']}").json()
    assert d["risk_class"] == "LOW"                                    # generated result untouched
    assert d["effective"]["risk_class"] == "MEDIUM" and d["effective"]["ai_usage"] == "YES"
    assert d["reviews"][0]["original_snapshot"]["risk_class"] == "LOW"


def test_reassessment_creates_new_run(client, completed):
    a = next(x for x in completed["run"]["assessments"] if x["vendor_ref"] == "V-003")
    r = client.post(f"/api/assessments/{a['id']}/reviews",
                    json={"decision": "REQUEST_REASSESSMENT", "analyst": "J. Analyst", "reason": "New press release"})
    assert r.status_code == 200 and r.json()["reassessment_run_id"]
    new = client.get(f"/api/runs/{r.json()['reassessment_run_id']}").json()
    assert new["parent_run_id"] == completed["run"]["id"] and new["total_vendors"] == 1
    old = client.get(f"/api/assessments/{a['id']}").json()
    assert old["review_status"] == "REASSESSMENT_REQUESTED" and old["status"] == "COMPLETED"


def test_approve_and_excel_export(client, completed, sample_bytes):
    a = next(x for x in completed["run"]["assessments"] if x["vendor_ref"] == "V-002")
    assert client.post(f"/api/assessments/{a['id']}/reviews",
                       json={"decision": "APPROVE", "analyst": "J. Analyst"}).status_code == 200
    r = client.get(f"/api/exports/excel?import_id={completed['import_id']}")
    assert r.status_code == 200
    out = load_workbook(io.BytesIO(r.content))
    orig = load_workbook(io.BytesIO(sample_bytes))
    ws, wo = out["Vendor Inventory"], orig["Vendor Inventory"]
    for c in range(1, wo.max_column + 1):                              # V-000 untouched
        assert ws.cell(4, c).value == wo.cell(4, c).value
    for row in range(5, 11):                                           # provided columns untouched
        for c in range(2, 12):
            assert ws.cell(row, c).value == wo.cell(row, c).value
    assert ws.cell(6, 12).value == "Critical" and ws.cell(6, 15).value == "Yes"
    assert "reviewed by J. Analyst (approved" in ws.cell(6, 22).value
    assert ws.cell(8, 19).value == "Medium"                            # V-004 override applied
    assert ws.cell(10, 19).value == "None identified"                  # V-006
    assert "NO MEANINGFUL PUBLIC EVIDENCE" in ws.cell(10, 16).value
    assert {"Evidence Log", "Sources Checked", "Methodology"} <= set(out.sheetnames)
    assert out["Evidence Log"].max_row > 5


def test_json_export_and_audit(client, completed):
    r = client.get(f"/api/exports/json?import_id={completed['import_id']}")
    data = json.loads(r.content)
    assert len(data["assessments"]) == 6
    events = {e["event_type"] for e in client.get("/api/audit").json()}
    assert {"IMPORT_ACCEPTED", "RUN_CREATED", "RUN_STARTED", "VENDOR_ASSESSED", "RUN_COMPLETED",
            "REVIEW_OVERRIDE", "REVIEW_APPROVE", "EXPORT_EXCEL"} <= events


def test_dashboard(client, completed):
    d = client.get("/api/dashboard").json()
    assert d["totals"]["vendors"] == 6 and d["totals"]["assessed"] == 6
    assert d["totals"]["ai_confirmed"] >= 2


def test_methodology_versioning(client):
    m = client.get("/api/methodology").json()
    bad = client.put("/api/methodology", json={"yaml": m["active"]["yaml"].replace("operational_dependency: 0.30", "operational_dependency: 0.90"),
                                               "note": "bad weights"})
    assert bad.status_code == 400
    new_yaml = m["active"]["yaml"].replace("CRITICAL: 0.75\n    HIGH: 0.55", "CRITICAL: 0.80\n    HIGH: 0.55")
    ok = client.put("/api/methodology", json={"yaml": new_yaml, "note": "Raise Critical threshold"})
    assert ok.status_code == 200
    after = client.get("/api/methodology").json()
    assert after["active"]["id"] != m["active"]["id"] and len(after["versions"]) >= 2
    # restore original so other tests are unaffected
    assert client.put("/api/methodology", json={"yaml": m["active"]["yaml"], "note": "restore"}).status_code == 200


def test_one_failing_source_does_not_fail_vendor(db, sample_bytes, monkeypatch):
    from app.services.osint.providers import CachedSnapshotProvider, SearchError
    from app.services.runs import create_import, start_run
    from app.core.config import get_settings

    class Flaky(CachedSnapshotProvider):
        calls = 0

        def search(self, query, **kw):
            Flaky.calls += 1
            if Flaky.calls % 2 == 0:
                raise SearchError("simulated outage")
            return super().search(query, **kw)

    imp, _ = create_import(db, "m.xlsx", sample_bytes)
    run = start_run(db, imp.id, vendor_refs=["V-002", "V-005"])
    orchestrator.run_assessment(run.id, provider=Flaky(get_settings().demo_cache_path))
    db.expire_all()
    from app.models import AssessmentRun
    r = db.get(AssessmentRun, run.id)
    assert r.completed_vendors == 2 and r.failed_vendors == 0 and r.status == "PARTIAL"
    assert all(a.errors for a in r.assessments)                        # failures are recorded, not hidden
