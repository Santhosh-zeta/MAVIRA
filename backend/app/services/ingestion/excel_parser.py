"""Meridian vendor-inventory workbook parser.

Columns are mapped by header text (not position), so the parser works on the
supplied workbook and on re-ordered copies. The fictional worked example (V-000)
is detected and preserved but never assessed or written to.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

# canonical field -> header aliases (normalised: lowercase, alphanumerics only)
INPUT_COLUMNS: dict[str, list[str]] = {
    "vendor_ref": ["vendorid"],
    "name": ["vendorname"],
    "description": ["vendordescription"],
    "service": ["serviceproductprovided", "serviceproductprovidedtomeridian"],
    "category": ["category"],
    "website": ["publicwebsite"],
    "business_process": ["businessprocesssupported"],
    "operational_dependency": ["operationaldependency"],
    "data_classification": ["dataclassificationaccessed"],
    "data_volume": ["datavolumeannual", "datavolume"],
}
OUTPUT_COLUMNS: dict[str, list[str]] = {
    "criticality_tier": ["criticalitytier"],
    "criticality_rationale": ["criticalityrationale"],
    "assessment_depth": ["assessmentdepthapplied"],
    "ai_usage": ["aiusagedetectedyn", "aiusagedetected"],
    "evidence": ["evidencedescriptionalongwithosintsourcesevidencelink", "evidencedescription"],
    "how_using_ai": ["howvendormaybeusingai"],
    "ai_providers": ["namedaisubprocessorsmodelproviders"],
    "risk_class": ["aisecurityriskclassification"],
    "risk_rationale": ["riskrationale"],
    "recommended_action": ["recommendedaction"],
    "assessed_by": ["assessedbydate", "assesseddbydate"],   # workbook header has a typo
}
REQUIRED_INPUT = ["vendor_ref", "name", "service", "category", "website", "operational_dependency",
                  "data_classification"]
VENDOR_REF_RE = re.compile(r"^V-\d{3,}$", re.I)


def norm_header(v: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())


@dataclass
class ParsedVendor:
    row_number: int
    is_example: bool
    fields: dict[str, str | None]
    raw: dict[str, Any]


@dataclass
class ParseResult:
    sheet_name: str
    header_row: int
    column_map: dict[str, int]
    detected_columns: list[str]
    vendors: list[ParsedVendor] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    missing_input_columns: list[str] = field(default_factory=list)
    missing_output_columns: list[str] = field(default_factory=list)
    duplicate_vendor_ids: list[str] = field(default_factory=list)

    @property
    def assessable(self) -> list[ParsedVendor]:
        return [v for v in self.vendors if not v.is_example]

    @property
    def valid(self) -> bool:
        return not self.errors

    def report(self) -> dict[str, Any]:
        return {
            "valid": self.valid,
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "detected_columns": self.detected_columns,
            "column_map": self.column_map,
            "vendor_count": len(self.assessable),
            "example_rows": [v.fields.get("vendor_ref") for v in self.vendors if v.is_example],
            "missing_input_columns": self.missing_input_columns,
            "missing_output_columns": self.missing_output_columns,
            "duplicate_vendor_ids": self.duplicate_vendor_ids,
            "errors": self.errors,
            "warnings": self.warnings,
            "vendors": [
                {"row": v.row_number, "vendor_ref": v.fields.get("vendor_ref"), "name": v.fields.get("name"),
                 "is_example": v.is_example,
                 "missing_fields": [k for k in REQUIRED_INPUT if not v.fields.get(k)]}
                for v in self.vendors
            ],
        }


class WorkbookError(ValueError):
    pass


def _find_header(ws) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        values = {norm_header(c.value) for c in row}
        if "vendorid" in values and "vendorname" in values:
            return row[0].row
    return None


def _match(aliases: dict[str, list[str]], headers: dict[str, int]) -> tuple[dict[str, int], list[str]]:
    found, missing = {}, []
    for key, opts in aliases.items():
        col = next((headers[o] for o in opts if o in headers), None)
        if col is None:   # tolerate small header variations via prefix match
            col = next((c for h, c in headers.items() if any(h.startswith(o[:12]) for o in opts)), None)
        if col is None:
            missing.append(key)
        else:
            found[key] = col
    return found, missing


def parse_workbook(data: bytes) -> ParseResult:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        raise WorkbookError(f"Workbook could not be opened: {type(exc).__name__}") from exc

    ws = next((wb[n] for n in wb.sheetnames if "vendor" in n.lower() and "inventory" in n.lower()), None)
    if ws is None:
        ws = next((s for s in wb.worksheets if _find_header(s)), None)
    if ws is None:
        raise WorkbookError("No sheet with a 'Vendor ID' / 'Vendor Name' header row was found")

    header_row = _find_header(ws)
    if header_row is None:
        raise WorkbookError(f"Sheet '{ws.title}' has no recognisable header row")

    headers: dict[str, int] = {}
    detected: list[str] = []
    for cell in ws[header_row]:
        if cell.value not in (None, ""):
            headers[norm_header(cell.value)] = cell.column
            detected.append(str(cell.value).strip())

    in_map, missing_in = _match(INPUT_COLUMNS, headers)
    out_map, missing_out = _match(OUTPUT_COLUMNS, headers)
    result = ParseResult(sheet_name=ws.title, header_row=header_row,
                         column_map={**in_map, **{f"out:{k}": v for k, v in out_map.items()}},
                         detected_columns=detected, missing_input_columns=missing_in,
                         missing_output_columns=missing_out)
    for k in missing_in:
        (result.errors if k in REQUIRED_INPUT else result.warnings).append(f"Input column not found: {k}")
    for k in missing_out:
        result.warnings.append(f"Assessment column not found (will be appended on export): {k}")
    if "vendor_ref" not in in_map:
        return result

    seen: dict[str, int] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        ref = ws.cell(r, in_map["vendor_ref"]).value
        if ref in (None, ""):
            continue
        ref = str(ref).strip()
        fields = {k: (str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else None)
                  for k, c in in_map.items()}
        name = fields.get("name") or ""
        is_example = ref.upper() == "V-000" or "example vendor" in name.lower() or "do not edit" in name.lower()
        if not VENDOR_REF_RE.match(ref):
            result.warnings.append(f"Row {r}: vendor id '{ref}' does not match V-### format")
        if ref in seen:
            result.duplicate_vendor_ids.append(ref)
            result.errors.append(f"Duplicate vendor id {ref} on rows {seen[ref]} and {r}")
            continue
        seen[ref] = r
        if not is_example:
            missing = [k for k in REQUIRED_INPUT if not fields.get(k)]
            if missing:
                result.warnings.append(f"{ref}: missing values for {', '.join(missing)}")
            prefilled = [k for k, c in out_map.items() if ws.cell(r, c).value not in (None, "")]
            if prefilled:
                result.warnings.append(f"{ref}: assessment columns already populated ({', '.join(prefilled)}); export will overwrite them")
        raw = {str(ws.cell(header_row, c).value): ws.cell(r, c).value for c in range(1, ws.max_column + 1)
               if ws.cell(header_row, c).value}
        result.vendors.append(ParsedVendor(row_number=r, is_example=is_example, fields=fields, raw=raw))

    if not result.assessable:
        result.errors.append("No assessable vendor rows found (only the worked example or none)")
    return result
